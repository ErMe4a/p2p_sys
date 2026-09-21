"""
Построчный импорт январских сделок из официальной "Книги учёта доходов и
расходов" (лист "Табл.№1") в реальные Order-записи.

Контекст (openspec/history.md §18): раньше форма "Загрузка данных из Excel"
(orders.views.admin_manual_entry_upload_excel) читала из этого же листа
только итоговые числа за месяц в MonthlyManualEntry, ни одна отдельная
сделка не сохранялась. Эта команда — вариант "1" из §18: настоящие Order
на каждую строку сделки.

Намеренно НЕ переиспользует api_views.order()/views-код создания ордера —
тот путь синхронно шлёт фискальный чек через receipt_service, что для
семи-восьмимесячных январских сделок не нужно (подтверждено пользователем).
Созданные ордера получают is_verified=True (числа уже подтверждены
официальным документом ИП, а не API биржи — но это тот же флаг, который
tasks.retry_blocked_orders_task использует как фильтр "уже не трогать";
намеренно проставляем его, чтобы этот периодический таск не подхватил
январские ордера и не попытался бы верифицировать/фискализировать их
через живой API биржи).

По умолчанию — dry-run: только парсит файл и сверяет построчную сумму с
итоговой строкой "Торговый результат" внутри самого файла. Ничего не
пишет в БД, пока не передан --commit.
"""
import datetime
import re
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from orders.models import Order

RECONCILE_TOLERANCE = 1.0  # ₽ / единицы — расхождение больше этого блокирует commit

# Определение биржи по формату "Номера документа" — в файле колонки "биржа"
# нет вообще, поэтому единственный источник — форма самого ID. Правила ниже
# проверены НЕ только по форме (совпадение с regex'ами извлечения orderId в
# content.js/mexc.js и т.п.), а подтверждены read-only запросом к проду:
# для каждой формы посчитано реальное распределение exchange_type среди уже
# существующих (Feb+) верифицированных Order с таким же форматом ID.
#   - ровно 19 цифр                      -> Bybit    (доминирующий формат)
#   - ровно 16 цифр                      -> Bitget    (348/348 = 100% в проде)
#   - ровно 8 цифр                       -> Gate      (702/730 = 96% в проде,
#                                            остаток — Bitget, статистический шум)
#   - латинская буква + 15-25 цифр       -> MEXC      (единственный формат
#                                            из "прочих", что проходит
#                                            regex MEXC-плашки \w{15,30})
#   - "OS-"/"OB-" + 10 алфанумерических  -> Telegram  (44025/44188 = 99.6%
#                                            в проде; S/B = Sell/Buy в id)
# Всё, что не подошло ни под одно правило — НЕ угадываем, помечаем явной
# меткой на ручную проверку.
_BYBIT_ID_RE    = re.compile(r'^\d{19}$')
_BITGET16_ID_RE = re.compile(r'^\d{16}$')
_GATE_ID_RE     = re.compile(r'^\d{8}$')
_MEXC_ID_RE     = re.compile(r'^[a-zA-Z]\d{15,25}$')
_TELEGRAM_ID_RE = re.compile(r'^O[SB]-\w{10}$')

UNKNOWN_EXCHANGE_LABEL = 'Импорт-уточнить биржу'


def _infer_exchange_type(doc_no: str) -> str:
    s = doc_no.strip()
    if _BYBIT_ID_RE.match(s):
        return 'Bybit'
    if _BITGET16_ID_RE.match(s):
        return 'Bitget'
    if _GATE_ID_RE.match(s):
        return 'Gate'
    if _MEXC_ID_RE.match(s):
        return 'MEXC'
    if _TELEGRAM_ID_RE.match(s):
        return 'Telegram'
    return UNKNOWN_EXCHANGE_LABEL


def _derive_username(file_path: str) -> str:
    import os
    return os.path.basename(file_path).rsplit('.', 1)[0].strip().lower()


def _find_sheet(wb):
    for name in wb.sheetnames:
        if 'табл' in name.lower() and '№1' in name.lower().replace(' ', ''):
            return wb[name]
    for name in wb.sheetnames:
        if 'табл' in name.lower() and '1' in name:
            return wb[name]
    return None


class Command(BaseCommand):
    help = (
        'Построчный импорт январских сделок из "Табл.№1" официальной книги учёта '
        'в реальные Order-записи. Без --commit — только парсинг и сверка (dry-run).'
    )

    def add_arguments(self, parser):
        parser.add_argument('file', help='Путь к .xlsx файлу (имя файла = username трейдера)')
        parser.add_argument('--username', help='Override username, если не совпадает с именем файла')
        parser.add_argument('--month', default='2026-01', help='Целевой месяц YYYY-MM — строки с датой вне него пропускаются (по умолчанию январь 2026)')
        parser.add_argument('--commit', action='store_true', help='Реально создать Order-записи (без флага — только dry-run)')

    def handle(self, *args, **options):
        path = options['file']
        username = (options.get('username') or _derive_username(path)).lower()
        commit = options['commit']
        target_year, target_month = (int(x) for x in options['month'].split('-'))

        User = get_user_model()
        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            raise CommandError(f'Пользователь "{username}" не найден в БД')

        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = _find_sheet(wb)
        if sheet is None:
            raise CommandError('Лист "Табл.№1" не найден в файле')

        rows = list(sheet.iter_rows(values_only=True))

        parsed = []
        summary = {}
        issues = []

        for row in rows:
            if not row or len(row) < 12:
                continue
            label = str(row[0] or '').strip().lower()

            if 'торговый результат' in label:
                summary['buy_qty']  = float(row[4] or 0)
                summary['buy_cost'] = float(row[6] or 0)
                summary['sell_qty']  = float(row[8] or 0)
                summary['sell_cost'] = float(row[10] or 0)
                continue

            if 'реализованный результат' in label or 'остаток' in label or label.startswith('прибыль'):
                continue

            num = row[0]
            if num is None or not str(num).strip().isdigit():
                continue  # заголовок/пустая строка

            date_raw, doc_no, currency = row[1], row[2], row[3]
            buy_qty, buy_price, buy_cost, buy_comm = row[4], row[5], row[6], row[7]
            sell_qty, sell_price, sell_cost, sell_comm = row[8], row[9], row[10], row[11]

            buy_cost_f  = float(buy_cost or 0)
            sell_cost_f = float(sell_cost or 0)

            if buy_cost_f and sell_cost_f:
                issues.append(f'строка №{num}: заполнены и покупка, и продажа одновременно — пропущена, нужен ручной разбор')
                continue
            if not buy_cost_f and not sell_cost_f:
                continue  # пустая строка без сделки

            if buy_cost_f:
                op_type, amount, price, cost, comm = 'BUY', buy_qty, buy_price, buy_cost, buy_comm
            else:
                op_type, amount, price, cost, comm = 'SELL', sell_qty, sell_price, sell_cost, sell_comm

            try:
                created_at = datetime.datetime.strptime(str(date_raw).strip(), '%d.%m.%Y').replace(hour=12)
            except (ValueError, TypeError):
                issues.append(f'строка №{num}: не удалось разобрать дату "{date_raw}" — пропущена')
                continue

            # Найдено на практике: некоторые "январские" файлы на деле продолжают
            # ту же таблицу в феврале/марте (трейдер не остановился на 31.01).
            # Такие строки — уже не January-only забой, они попадают в реальный
            # расчёт прибыли/налога (SYSTEM_START их не фильтрует) и никогда не
            # получат чек (is_verified=True сразу исключает их из
            # retry_blocked_orders_task) — поэтому жёстко режем, а не молча
            # тащим за пределы целевого месяца.
            if (created_at.year, created_at.month) != (target_year, target_month):
                issues.append(
                    f'строка №{num}: дата {date_raw} вне целевого месяца '
                    f'{target_year}-{target_month:02d} — пропущена (не январь)'
                )
                continue

            doc_no_str = str(doc_no).strip() if doc_no is not None else ''
            if not doc_no_str:
                issues.append(f'строка №{num}: пустой номер документа — пропущена')
                continue

            parsed.append({
                'row_num':        num,
                'external_id':    doc_no_str,
                'exchange_type':  _infer_exchange_type(doc_no_str),
                'currency':       str(currency).strip() if currency else 'USDT',
                'operation_type': op_type,
                'amount':         Decimal(str(amount or 0)),
                'price':          Decimal(str(price or 0)),
                'cost':           Decimal(str(cost or 0)),
                'commission':     Decimal(str(comm or 0)),
                'created_at':     timezone.make_aware(created_at),
            })

        buy_cost_sum  = sum((p['cost'] for p in parsed if p['operation_type'] == 'BUY'), Decimal('0'))
        sell_cost_sum = sum((p['cost'] for p in parsed if p['operation_type'] == 'SELL'), Decimal('0'))
        buy_qty_sum   = sum((p['amount'] for p in parsed if p['operation_type'] == 'BUY'), Decimal('0'))
        sell_qty_sum  = sum((p['amount'] for p in parsed if p['operation_type'] == 'SELL'), Decimal('0'))

        exchange_counts = {}
        for p in parsed:
            exchange_counts[p['exchange_type']] = exchange_counts.get(p['exchange_type'], 0) + 1

        self.stdout.write(f'Файл: {path}')
        self.stdout.write(f'Пользователь: {username} (id={user.id})')
        self.stdout.write(f'Строк-сделок разобрано: {len(parsed)}')
        self.stdout.write(f'Распределение по бирже (определено по формату номера документа): {exchange_counts}')
        if UNKNOWN_EXCHANGE_LABEL in exchange_counts:
            self.stdout.write(self.style.WARNING(
                f'  Внимание: {exchange_counts[UNKNOWN_EXCHANGE_LABEL]} строк с неопознанным форматом ID — '
                f'помечены "{UNKNOWN_EXCHANGE_LABEL}", биржа не проставлена, требуется уточнение у Макса.'
            ))

        if issues:
            self.stdout.write(self.style.WARNING(f'Проблемные строки, пропущены ({len(issues)}):'))
            for i in issues:
                self.stdout.write(f'  - {i}')

        self.stdout.write('Сверка построчной суммы с итоговой строкой "Торговый результат" внутри файла:')
        self.stdout.write(
            f'  Покупки: построчно cost={buy_cost_sum:.2f} qty={buy_qty_sum:.6f}  |  '
            f'в файле cost={summary.get("buy_cost", 0):.2f} qty={summary.get("buy_qty", 0):.6f}'
        )
        self.stdout.write(
            f'  Продажи: построчно cost={sell_cost_sum:.2f} qty={sell_qty_sum:.6f}  |  '
            f'в файле cost={summary.get("sell_cost", 0):.2f} qty={summary.get("sell_qty", 0):.6f}'
        )

        mismatch = (
            abs(float(buy_cost_sum) - summary.get('buy_cost', 0)) > RECONCILE_TOLERANCE
            or abs(float(sell_cost_sum) - summary.get('sell_cost', 0)) > RECONCILE_TOLERANCE
            or abs(float(buy_qty_sum) - summary.get('buy_qty', 0)) > RECONCILE_TOLERANCE
            or abs(float(sell_qty_sum) - summary.get('sell_qty', 0)) > RECONCILE_TOLERANCE
        )
        if mismatch:
            self.stdout.write(self.style.ERROR(
                'РАСХОЖДЕНИЕ построчной суммы с итоговой строкой файла — commit заблокирован, нужен ручной разбор.'
            ))
            return

        self.stdout.write(self.style.SUCCESS('Сверка сошлась — построчные суммы совпадают с итогом файла.'))

        if not commit:
            self.stdout.write(self.style.WARNING(
                'Dry-run: в БД ничего не записано. Запусти с --commit для реальной записи.'
            ))
            return

        existing_ids = set(
            Order.objects.filter(user=user, external_id__in=[p['external_id'] for p in parsed])
            .values_list('external_id', flat=True)
        )

        created = 0
        skipped = 0
        for p in parsed:
            if p['external_id'] in existing_ids:
                skipped += 1
                continue
            Order.objects.create(
                user=user,
                external_id=p['external_id'],
                price=p['price'],
                amount=p['amount'],
                cost=p['cost'],
                operation_type=p['operation_type'],
                exchange_type=p['exchange_type'],
                currency=p['currency'],
                commission=p['commission'],
                commission_type='FIX',
                created_at=p['created_at'],
                is_verified=True,
                is_manual=True,
            )
            created += 1

        self.stdout.write(self.style.SUCCESS(f'Создано Order: {created}, пропущено (уже существовал external_id): {skipped}'))
